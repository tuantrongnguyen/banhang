"""Xuất Excel: hóa đơn 1 đơn hàng và danh sách đơn hàng (kèm chi tiết mặt hàng)."""
import math
from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VN = timedelta(hours=7)  # DB lưu UTC, file Excel hiển thị giờ Việt Nam
PAY = {"cash": "Tiền mặt", "transfer": "Chuyển khoản", "card": "Quẹt thẻ"}

FONT = "Arial"
MONEY = "#,##0"
DATE = "dd/mm/yyyy hh:mm"
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HEAD_FILL = PatternFill("solid", fgColor="13213B")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Border(top=Side(style="medium", color="13213B"))
TOTAL_FILL = PatternFill("solid", fgColor="F3F2EE")


def local(d):
    return d + VN


def line_net(price, qty, pct):
    return int(math.floor(price * qty * (100 - (pct or 0)) / 100 + 0.5))


def _header(ws, row, titles):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _put(ws, row, col, value, fmt=None, font=BODY, align=None, border=BOX):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    if border:
        cell.border = border
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = Alignment(horizontal=align, vertical="center")
    return cell


def _save(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_orders_xlsx(orders, items):
    """orders: list[Order]; items: list[OrderItem] (chỉ lấy những dòng thuộc `orders`)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Đơn hàng"
    heads = ["Mã đơn", "Thời gian", "Khách hàng", "Thanh toán", "Tạm tính (sau CK mặt hàng)",
             "Chiết khấu cả đơn", "Phí ship", "Tổng cộng", "Lợi nhuận", "Ghi chú"]
    _header(ws, 1, heads)
    for r, o in enumerate(orders, 2):
        _put(ws, r, 1, o.id, align="center")
        _put(ws, r, 2, local(o.created_at), DATE, align="center")
        _put(ws, r, 3, o.customer_name)
        _put(ws, r, 4, PAY.get(o.payment, o.payment))
        _put(ws, r, 5, o.subtotal, MONEY)
        _put(ws, r, 6, o.discount, MONEY)
        _put(ws, r, 7, o.shipping, MONEY)
        _put(ws, r, 8, o.total, MONEY)
        _put(ws, r, 9, o.profit, MONEY)
        _put(ws, r, 10, o.note or "")
    last = len(orders) + 1
    tr = last + 1
    _put(ws, tr, 1, "TỔNG", font=BOLD, border=TOP)
    for c in range(2, 5):
        _put(ws, tr, c, None, border=TOP)
    for c in range(5, 10):
        L = get_column_letter(c)
        cell = _put(ws, tr, c, f"=SUM({L}2:{L}{last})" if orders else 0, MONEY, BOLD, border=TOP)
        cell.fill = TOTAL_FILL
    _put(ws, tr, 10, None, border=TOP)
    ws.cell(row=tr, column=1).fill = TOTAL_FILL
    _widths(ws, [10, 18, 26, 16, 20, 18, 14, 16, 16, 28])
    ws.freeze_panes = "A2"
    if orders:
        ws.auto_filter.ref = f"A1:J{last}"

    # --- chi tiết mặt hàng ---
    ws2 = wb.create_sheet("Chi tiết mặt hàng")
    _header(ws2, 1, ["Mã đơn", "Thời gian", "Sản phẩm", "Số lượng", "Đơn giá", "Chiết khấu (%)", "Thành tiền", "Giá vốn"])
    by_id = {o.id: o for o in orders}
    r = 2
    for it in items:
        o = by_id.get(it.order_id)
        if not o:
            continue
        _put(ws2, r, 1, it.order_id, align="center")
        _put(ws2, r, 2, local(o.created_at), DATE, align="center")
        _put(ws2, r, 3, it.product_name)
        _put(ws2, r, 4, it.qty, "0", align="center")
        _put(ws2, r, 5, it.price, MONEY)
        _put(ws2, r, 6, it.discount_pct or 0, "0.##", align="center")
        _put(ws2, r, 7, line_net(it.price, it.qty, it.discount_pct), MONEY)
        _put(ws2, r, 8, it.cost * it.qty, MONEY)
        r += 1
    _widths(ws2, [10, 18, 34, 10, 16, 14, 16, 16])
    ws2.freeze_panes = "A2"
    if r > 2:
        ws2.auto_filter.ref = f"A1:H{r - 1}"
    return _save(wb)


def build_invoice_xlsx(order, items, draft=False):
    """Hóa đơn của 1 đơn hàng. Số tiền lấy nguyên từ database (đã làm tròn theo quy tắc của hệ thống)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Phiếu tạm tính" if draft else "Hóa đơn"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "SỔ BÁN HÀNG"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color="13213B")
    ws["A2"] = "PHIẾU TẠM TÍNH (CHƯA THANH TOÁN)" if draft else f"HÓA ĐƠN BÁN HÀNG #{order.id}"
    ws["A2"].font = Font(name=FONT, size=11, bold=True)
    info = [("Ngày", local(order.created_at)), ("Khách hàng", order.customer_name), ("Thanh toán", PAY.get(order.payment, order.payment))]
    for i, (k, v) in enumerate(info, 4):
        ws.cell(row=i, column=1, value=k).font = BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font = BODY
        c.alignment = Alignment(horizontal="left")
        if k == "Ngày":
            c.number_format = DATE

    head_row = 8
    _header(ws, head_row, ["STT", "Sản phẩm", "SL", "Đơn giá", "CK (%)", "Thành tiền"])
    r = head_row + 1
    gross = 0
    for n, it in enumerate(items, 1):
        gross += it.price * it.qty
        _put(ws, r, 1, n, align="center")
        _put(ws, r, 2, it.product_name)
        _put(ws, r, 3, it.qty, "0", align="center")
        _put(ws, r, 4, it.price, MONEY)
        _put(ws, r, 5, it.discount_pct or 0, "0.##", align="center")
        _put(ws, r, 6, line_net(it.price, it.qty, it.discount_pct), MONEY)
        r += 1

    item_disc = max(0, gross - order.subtotal)
    pct = (order.discount / order.subtotal * 100) if order.subtotal else 0
    r += 1
    rows = [("Tạm tính", gross, BODY)]
    if item_disc > 0:
        rows.append(("Chiết khấu mặt hàng", -item_disc, BODY))
    if order.discount > 0:
        rows.append((f"Chiết khấu cả đơn ({pct:.1f}%)".replace(".0%", "%"), -order.discount, BODY))
    rows += [("Sau chiết khấu", order.subtotal - order.discount, BODY), ("Phí ship", order.shipping, BODY)]
    for label, val, f in rows:
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.cell(row=r, column=3).value = label
        ws.cell(row=r, column=3).font = f
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="right")
        _put(ws, r, 6, val, MONEY, f, border=None)
        r += 1
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    t = ws.cell(row=r, column=3, value="TỔNG CỘNG")
    t.font = Font(name=FONT, size=12, bold=True)
    t.alignment = Alignment(horizontal="right")
    tv = ws.cell(row=r, column=6, value=order.total)
    tv.font, tv.number_format, tv.fill = Font(name=FONT, size=12, bold=True), MONEY, TOTAL_FILL
    r += 2
    ws.cell(row=r, column=1, value="Cảm ơn quý khách!").font = Font(name=FONT, size=10, italic=True)
    _widths(ws, [14, 38, 8, 16, 10, 18])
    return _save(wb)
